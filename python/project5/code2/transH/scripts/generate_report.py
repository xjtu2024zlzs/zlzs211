from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from services.pipe_design_reference import apply_pipe_design_reference

STAGE_ADVICE = {
    "design": "建议反馈设计子系统，复核设计参数、裕度、结构约束。",
    "material": "建议反馈材料/供应商管理子系统，复核材料批次、硬度、疲劳性能。",
    "manufacturing": "建议反馈制造工艺子系统，复核加工误差、表面粗糙度、洁净度、热处理或磨削过程。",
    "assembly": "建议反馈装配子系统，复核装配扭矩、预紧、对中、装配人员和工位。",
    "inspection": "建议反馈检测子系统，复核检测设备、检测记录和异常指标。",
    "operation": "建议反馈运维保障子系统，复核维护、润滑、污染、老化和使用状态。",
}

STAGE_SUBSYSTEMS_READABLE = {
    "design": "设计子系统",
    "material": "材料/供应商管理子系统",
    "manufacturing": "制造工艺子系统",
    "assembly": "装配子系统",
    "inspection": "检测子系统",
    "operation": "运维保障子系统",
    "unknown": "待人工复核子系统",
}

HYDRAULIC_NAMES = {
    "C001": "液压泵总成",
    "C002": "方向控制阀",
    "C003": "作动筒",
    "C004": "溢流阀",
    "C005": "蓄能器",
    "C006": "冷却器",
    "C007": "过滤器",
    "C008": "液压油箱",
    "C009": "单向阀",
    "C010": "节流阀",
    "C011": "管路总成",
    "C012": "卸荷阀",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate lifecycle quality traceability reports.")
    parser.add_argument("--domain", choices=["bearing", "hydraulic"], required=True)
    parser.add_argument("--case_id", required=True)
    parser.add_argument("--print_json", action="store_true", help="Print generated report JSON payload.")
    parser.add_argument("--no_docx", action="store_true", help="Skip Word docx generation.")
    return parser.parse_args()


def sanitize_text(value: Any) -> str:
    return (
        str(value)
        .replace("调速阀/节流阀", "节流阀")
        .replace("密封件/管路总成", "管路总成")
        .replace("同时考虑 了", "同时考虑了")
    )


def normalize_case_id_alias(case_id: str) -> str:
    bearing_match = re.fullmatch(r"FB-BRG-(\d+)", case_id)
    if bearing_match:
        return f"BFB-{int(bearing_match.group(1)):05d}"
    hydraulic_match = re.fullmatch(r"FB-HYD-(\d+)", case_id)
    if hydraulic_match:
        return f"HFB-{int(hydraulic_match.group(1)):05d}"
    return case_id


def require_file(path: Path, hint: str) -> Path:
    if not path.exists():
        raise FileNotFoundError(f"缺少文件: {path}\n请先完成: {hint}")
    return path


def latest_reasoning_path(domain: str) -> Path:
    report_dir = ROOT / "outputs" / domain / "reports"
    files = sorted(report_dir.glob("reasoning_*.json"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        raise FileNotFoundError(f"{report_dir} 下没有 reasoning_*.json，请先运行 scripts/run_reasoning.py。")
    return files[0]


def resolve_paths(domain: str, case_id_arg: str) -> dict[str, Path]:
    report_dir = ROOT / "outputs" / domain / "reports"
    graph_dir = ROOT / "outputs" / domain / "graphs"
    if case_id_arg == "latest":
        reasoning = latest_reasoning_path(domain)
        case_id = reasoning.stem.replace("reasoning_", "", 1)
    else:
        case_id = normalize_case_id_alias(case_id_arg)
        reasoning = report_dir / f"reasoning_{case_id}.json"
    paths = {
        "reasoning": require_file(reasoning, f"python scripts/run_reasoning.py --domain {domain} --case_id {case_id_arg}"),
        "original_graph": require_file(graph_dir / f"original_graph_{case_id}.json", f"python scripts/export_graph.py --domain {domain} --case_id {case_id_arg}"),
        "fault_graph": require_file(graph_dir / f"fault_enhanced_graph_{case_id}.json", f"python scripts/export_graph.py --domain {domain} --case_id {case_id_arg}"),
        "evidence": ROOT / "data" / domain / f"{domain}_evidence.csv",
        "kg_statistics": report_dir / "kg_statistics.json",
        "train_log": report_dir / "train_log.csv",
    }
    if domain == "hydraulic":
        hdir = ROOT / "data" / "hydraulic"
        paths.update(
            {
                "rca": hdir / "hydraulic_rca_case.csv",
                "diagnosis": hdir / "hydraulic_component_diagnosis.csv",
                "subtype": hdir / "hydraulic_subtype_probability.csv",
                "sensor": hdir / "hydraulic_sensor_energy.csv",
            }
        )
    return paths


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def optional_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, encoding="utf-8-sig") if path.exists() else pd.DataFrame()


def optional_json(path: Path) -> dict[str, Any]:
    return load_json(path) if path.exists() else {}


def graph_stage_summary(graph: dict[str, Any]) -> list[str]:
    stages = []
    for node in graph.get("nodes", []):
        stage = node.get("stage")
        stage_name = node.get("stage_name")
        if stage and stage not in {"unknown", "feedback", "fault", "rca"} and stage_name not in stages:
            stages.append(stage_name)
    return stages


def target_lifecycle_summary(original_graph: dict[str, Any]) -> list[dict[str, Any]]:
    target_id = original_graph.get("summary", {}).get("target_object")
    rows = []
    node_map = {node["id"]: node for node in original_graph.get("nodes", [])}
    for edge in original_graph.get("edges", []):
        if edge["source"] == target_id:
            target_node = node_map.get(edge["target"], {})
            rows.append(
                {
                    "relation": edge["relation"],
                    "relation_name": edge.get("relation_name", edge["relation"]),
                    "node_id": edge["target"],
                    "node_name": target_node.get("name", edge["target"]),
                    "node_type": target_node.get("type"),
                    "stage_name": target_node.get("stage_name"),
                }
            )
    return rows


def graph_comparison(original_graph: dict[str, Any], fault_graph: dict[str, Any]) -> dict[str, Any]:
    original_nodes = {node["id"] for node in original_graph.get("nodes", [])}
    fault_nodes = {node["id"] for node in fault_graph.get("nodes", [])}
    added_nodes = [node for node in fault_graph.get("nodes", []) if node["id"] not in original_nodes]
    added_types = sorted({node.get("type", "unknown") for node in added_nodes})
    return {
        "original_graph": {
            "node_count": original_graph["summary"]["node_count"],
            "edge_count": original_graph["summary"]["edge_count"],
            "stages": graph_stage_summary(original_graph),
            "target_lifecycle_nodes": target_lifecycle_summary(original_graph),
        },
        "fault_enhanced_graph": {
            "node_count": fault_graph["summary"]["node_count"],
            "edge_count": fault_graph["summary"]["edge_count"],
            "added_node_count": len(fault_nodes - original_nodes),
            "added_edge_count": fault_graph["summary"]["edge_count"] - original_graph["summary"]["edge_count"],
            "added_node_types": added_types,
        },
        "difference_description": f"故障增强图在原始生命周期链基础上新增 {len(fault_nodes - original_nodes)} 个节点，主要包括 {', '.join(added_types[:8])} 等，用于承载故障、RCA、Top6候选和阶段归因信息。",
    }


def evidence_summary(evidence: pd.DataFrame, feedback_id: str) -> dict[str, Any]:
    if evidence.empty:
        return {"rca_text": "", "fault_evolution_text": ""}
    rows = evidence[evidence["feedback_id"].astype(str) == feedback_id]
    result = {"rca_text": "", "fault_evolution_text": ""}
    for row in rows.itertuples(index=False):
        if str(row.evidence_type) == "rca_text":
            result["rca_text"] = sanitize_text(row.evidence_text)
        elif str(row.evidence_type) == "fault_evolution_text":
            result["fault_evolution_text"] = sanitize_text(row.evidence_text)
    return result


def hydraulic_extra_summary(paths: dict[str, Path], reasoning: dict[str, Any]) -> dict[str, Any]:
    case_id = reasoning.get("rca_context", {}).get("case_id")
    ctx = dict(reasoning.get("rca_context", {}))
    if not case_id:
        return ctx
    for key in ["rca", "diagnosis", "subtype", "sensor"]:
        require_file(paths[key], f"python scripts/generate_hydraulic_data.py")
    diagnosis = pd.read_csv(paths["diagnosis"], encoding="utf-8-sig")
    subtype = pd.read_csv(paths["subtype"], encoding="utf-8-sig")
    sensor = pd.read_csv(paths["sensor"], encoding="utf-8-sig")
    diag_top5 = diagnosis[diagnosis["case_id"].astype(str) == str(case_id)].sort_values("rank_no").head(5)
    sub_top5 = subtype[subtype["case_id"].astype(str) == str(case_id)].sort_values("probability", ascending=False).head(5)
    sensor_row = sensor[sensor["case_id"].astype(str) == str(case_id)]
    sensor_top5: dict[str, float] = {}
    if not sensor_row.empty:
        values = sensor_row.iloc[0].drop(labels=["case_id"]).astype(float)
        sensor_top5 = {str(k): float(v) for k, v in values.sort_values(ascending=False).head(5).items()}
    ctx["component_diagnosis_top5"] = sanitize_records(diag_top5.to_dict(orient="records"))
    ctx["subtype_probability_top5"] = sanitize_records(sub_top5.to_dict(orient="records"))
    ctx["sensor_energy_top5"] = sensor_top5
    root_code = str(ctx.get("root_component_code", ""))
    if root_code in HYDRAULIC_NAMES:
        ctx["root_component_name"] = HYDRAULIC_NAMES[root_code]
    feedback_component = reasoning.get("target_object", {}).get("object_name", "")
    ctx["propagation_note"] = ""
    if root_code and sanitize_text(ctx.get("root_component_name", "")) not in sanitize_text(feedback_component):
        ctx["propagation_note"] = "RCA根因部件与当前反馈部件不完全一致，可能表示故障传播或相邻部件诱发。"
    return ctx


def sanitize_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    clean = []
    for record in records:
        clean.append({key: sanitize_text(value) if isinstance(value, str) else value for key, value in record.items()})
    return clean


def build_evidence_chains(reasoning: dict[str, Any], original_graph: dict[str, Any], fault_graph: dict[str, Any]) -> list[dict[str, Any]]:
    feedback_id = reasoning["feedback_summary"]["feedback_id"]
    target_id = reasoning["target_object"]["object_id"]
    top1 = reasoning["top6_candidates"][0]
    chains = [
        {
            "chain": [feedback_id, "occurs_on", target_id],
            "description": "质量反馈首先定位到发生故障的目标对象。",
        }
    ]
    preferred = [item for item in reasoning["top6_candidates"] if item["candidate_scope"] == "lifecycle_candidate"]
    for item in (preferred[:2] or reasoning["top6_candidates"][:2]):
        relation = find_relation_to_candidate(original_graph, target_id, item["candidate_id"])
        chains.append(
            {
                "chain": [target_id, relation or "lifecycle_relation", item["candidate_id"]],
                "description": f"候选 {item['candidate_id']} 位于目标对象生命周期链或其关键扩展路径上，path_score={item['path_score']}。",
            }
        )
    chains.append(
        {
            "chain": [feedback_id, "may_caused_by", top1["candidate_id"]],
            "description": f"TransH 综合重排序将 {top1['candidate_id']} 排为 Top1，final_score={top1['final_score']}。",
        }
    )
    return chains


def find_relation_to_candidate(graph: dict[str, Any], source: str, candidate_id: str) -> str | None:
    for edge in graph.get("edges", []):
        if edge["source"] == source and edge["target"] == candidate_id:
            return edge["relation"]
    return None


def explain_primary_stage(stage_attribution: dict[str, Any], top6: list[dict[str, Any]]) -> str:
    primary = stage_attribution["primary_stage"]
    primary_name = stage_attribution["primary_stage_name"]
    related = [item for item in top6 if item["candidate_stage"] == primary]
    names = "、".join(item["candidate_id"] for item in related[:3])
    return f"{primary_name}得分最高，主要因为 Top6 中 {names or '多个候选原因'} 聚集在该阶段，且综合分数、路径分数或证据分数较高。"


def normalize_focus_stage_top3(stage_attribution: dict[str, Any]) -> list[dict[str, Any]]:
    focus_items = stage_attribution.get("focus_stage_top3", [])
    if not isinstance(focus_items, list):
        return []
    rows = []
    for idx, item in enumerate(focus_items[:3], start=1):
        if not isinstance(item, dict):
            continue
        rows.append(
            {
                "rank": idx,
                "stage": str(item.get("stage", "")),
                "stage_name": str(item.get("stage_name", "")),
                "score": item.get("score", ""),
                "feedback_subsystem": str(item.get("feedback_subsystem", "")),
            }
        )
    return rows


def normalize_feedback_subsystem_top3(stage_attribution: dict[str, Any], focus_stage_top3: list[dict[str, Any]]) -> list[str]:
    subsystem_top3 = stage_attribution.get("feedback_subsystem_top3", [])
    if isinstance(subsystem_top3, list):
        result = [str(item) for item in subsystem_top3[:3] if str(item).strip()]
    else:
        result = []
    if not result:
        result = [row["feedback_subsystem"] for row in focus_stage_top3 if row.get("feedback_subsystem")]
    return result[:3]


def review_subsystem_candidates(stage_attribution: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for item in stage_attribution.get("stage_rank", []):
        if not isinstance(item, dict):
            continue
        stage = str(item.get("stage", "unknown"))
        rows.append(
            {
                "stage": stage,
                "stage_name": str(item.get("stage_name", "")),
                "score": item.get("score", ""),
                "feedback_subsystem": STAGE_SUBSYSTEMS_READABLE.get(stage, "待人工复核子系统"),
            }
        )
    return rows


def stage_top3_conclusion_note(stage_attribution: dict[str, Any]) -> str:
    focus_stage_top3 = normalize_focus_stage_top3(stage_attribution)
    subsystem_top3 = normalize_feedback_subsystem_top3(stage_attribution, focus_stage_top3)
    if not focus_stage_top3 and not subsystem_top3:
        return ""

    primary_subsystem = str(stage_attribution.get("feedback_target_subsystem", ""))
    joint_subsystems = [item for item in subsystem_top3 if item and item != primary_subsystem]
    if primary_subsystem and joint_subsystems:
        return (
            "除首要疑似生命周期阶段外，系统同时给出重点关注阶段 Top3，用于支持跨阶段联合复核。"
            f"建议优先反馈至{primary_subsystem}，并联合{'、'.join(joint_subsystems)}进行辅助复核。"
        )
    if primary_subsystem:
        return (
            "除首要疑似生命周期阶段外，系统同时给出重点关注阶段 Top3，用于支持跨阶段联合复核。"
            f"建议优先反馈至{primary_subsystem}，并结合重点关注阶段进行辅助复核。"
        )
    return "除首要疑似生命周期阶段外，系统同时给出重点关注阶段 Top3，用于支持跨阶段联合复核。"


def final_conclusion(reasoning: dict[str, Any], evidence_chains: list[dict[str, Any]], hydraulic_note: str = "") -> str:
    pipe_reference = reasoning.get("pipe_design_reference")
    reason_list = reasoning.get("reasonList")
    if isinstance(pipe_reference, dict) and pipe_reference.get("enabled") is True and isinstance(reason_list, list) and reason_list:
        top_reason = reason_list[0] if isinstance(reason_list[0], dict) else {}
        subsystem = reasoning["stage_attribution"]["feedback_target_subsystem"]
        evidence_text = str(top_reason.get("evidence") or "系统识别到C011管路总成与管路设计参数存在关联。")
        return (
            f"本次追溯的首要疑似原因是 {top_reason.get('reasonName', '管路设计参数复核')}，"
            "首要疑似阶段为设计阶段。"
            f"主要证据为：{evidence_text}"
            f"建议优先由{subsystem}牵头处置，并将管段长度、弯曲半径和弯曲角等参数提供给下游管路参数优化模块读取。"
        )

    top1 = reasoning["top6_candidates"][0]
    stage = reasoning["stage_attribution"]["primary_stage_name"]
    subsystem = reasoning["stage_attribution"]["feedback_target_subsystem"]
    evidence_text = evidence_chains[-1]["description"] if evidence_chains else "Top6 推理结果形成主要证据。"
    cross_note = hydraulic_note or "当前报告未发现明确跨部件传播描述，仍建议结合现场记录复核跨阶段影响。"
    top3_note = stage_top3_conclusion_note(reasoning.get("stage_attribution", {}))
    return (
        f"本次追溯的首要疑似原因是 {top1['candidate_id']}（{top1['candidate_type']}），首要疑似阶段为{stage}。"
        f"主要证据包括生命周期路径、TransH 评分、RCA/证据匹配以及规则先验，其中 {evidence_text}"
        f"{cross_note} 建议优先由{subsystem}牵头处置，并联动相关生命周期环节复核。{top3_note}"
    )


def build_report(domain: str, paths: dict[str, Path]) -> dict[str, Any]:
    reasoning = apply_pipe_design_reference(load_json(paths["reasoning"]))
    original_graph = load_json(paths["original_graph"])
    fault_graph = load_json(paths["fault_graph"])
    evidence = optional_csv(paths["evidence"])
    kg_stats = optional_json(paths["kg_statistics"])
    train_log = optional_csv(paths["train_log"])
    feedback_id = reasoning["feedback_summary"]["feedback_id"]
    ev = evidence_summary(evidence, feedback_id)
    rca_summary = ev
    hydraulic_note = ""
    if domain == "hydraulic":
        rca_summary = hydraulic_extra_summary(paths, reasoning)
        hydraulic_note = rca_summary.get("propagation_note", "")
    graph_cmp = graph_comparison(original_graph, fault_graph)
    evidence_chains = build_evidence_chains(reasoning, original_graph, fault_graph)
    stage = reasoning["stage_attribution"]
    focus_stage_top3 = normalize_focus_stage_top3(stage)
    feedback_subsystem_top3 = normalize_feedback_subsystem_top3(stage, focus_stage_top3)
    top6 = reasoning["top6_candidates"]
    report = {
        "report_meta": {
            "domain": domain,
            "case_id": reasoning["case_id"],
            "feedback_id": feedback_id,
            "report_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "model_name": "TransH",
            "top_k": 6,
            "source_files": {
                "reasoning_json": str(paths["reasoning"]),
                "original_graph_json": str(paths["original_graph"]),
                "fault_enhanced_graph_json": str(paths["fault_graph"]),
            },
            "kg_statistics": kg_stats,
            "train_log_last": train_log.tail(1).to_dict(orient="records")[0] if not train_log.empty else {},
        },
        "feedback_summary": reasoning["feedback_summary"],
        "target_object": reasoning["target_object"],
        "rca_summary": rca_summary,
        "graph_comparison": graph_cmp,
        "top6_reason_analysis": top6,
        "stage_attribution": {
            **stage,
            "focus_stage_top3": focus_stage_top3,
            "feedback_subsystem_top3": feedback_subsystem_top3,
            "primary_stage_explanation": explain_primary_stage(stage, top6),
        },
        "evidence_chains": evidence_chains,
        "subsystem_feedback": {
            "primary_stage": stage["primary_stage"],
            "target_subsystem": stage["feedback_target_subsystem"],
            "primary_feedback_subsystem": stage["feedback_target_subsystem"],
            "feedback_subsystem_top3": feedback_subsystem_top3,
            "review_subsystem_candidates": review_subsystem_candidates(stage),
            "suggestion": STAGE_ADVICE.get(stage["primary_stage"], "建议转入人工复核流程。"),
        },
        "final_conclusion": final_conclusion(reasoning, evidence_chains, hydraulic_note),
    }
    if isinstance(reasoning.get("pipe_design_reference"), dict):
        report["pipe_design_reference"] = reasoning["pipe_design_reference"]
    if isinstance(reasoning.get("reasonList"), list):
        report["reasonList"] = reasoning["reasonList"]
    return json.loads(sanitize_text(json.dumps(report, ensure_ascii=False)))


def md_table(headers: list[str], rows: list[list[Any]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(sanitize_text(x) for x in row) + " |")
    return "\n".join(lines)


def pipe_design_reference_enabled(report: dict[str, Any]) -> bool:
    reference = report.get("pipe_design_reference")
    return isinstance(reference, dict) and reference.get("enabled") is True


def priority_label(priority: Any) -> str:
    return {"high": "高", "medium": "中", "low": "低"}.get(str(priority), str(priority))


def pipe_reference_rows(reference: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for item in reference.get("parameter_review_items", []):
        if not isinstance(item, dict):
            continue
        rows.append(
            [
                item.get("param_name", ""),
                item.get("param_category", ""),
                priority_label(item.get("priority", "")),
                item.get("problem_type", ""),
                item.get("reason", ""),
                item.get("optimization_hint", ""),
            ]
        )
    return rows


def reason_list_rows(report: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for item in report.get("reasonList", []):
        if not isinstance(item, dict):
            continue
        rows.append(
            [
                item.get("rank", ""),
                item.get("reasonType", ""),
                item.get("relatedPart", ""),
                item.get("reasonName", ""),
                item.get("confidence", ""),
                item.get("evidence", ""),
                item.get("suggestion", ""),
            ]
        )
    return rows


def reason_ranking_markdown_lines(report: dict[str, Any]) -> list[str]:
    if pipe_design_reference_enabled(report) and isinstance(report.get("reasonList"), list):
        return [
            "## 6. 最终溯源原因表",
            "综合置信度综合考虑模型推理结果、图谱路径证据、反馈对象与生命周期阶段之间的关联关系以及管路设计参数证据。",
            md_table(
                ["排名", "原因阶段/类型", "原因名称", "关联对象", "置信度", "关键证据", "建议措施"],
                reason_list_rows(report),
            ),
        ]

    return [
        "## 6. TransH Top6疑似原因分析",
        md_table(
            ["rank", "candidate_id", "candidate_type", "stage", "final", "transh", "path", "evidence", "prior", "scope", "description"],
            [[c["rank"], c["candidate_id"], c["candidate_type"], c["candidate_stage_name"], c["final_score"], c["transh_score"], c["path_score"], c["evidence_score"], c["rule_prior_score"], c["candidate_scope"], c["reason_description"]] for c in report["top6_reason_analysis"]],
        ),
    ]


def pipe_reference_markdown_lines(report: dict[str, Any]) -> list[str]:
    if not pipe_design_reference_enabled(report):
        return []
    reference = report["pipe_design_reference"]
    variables = reference.get("optimization_variables", [])
    if not isinstance(variables, list):
        variables = []
    return [
        "",
        "## 11. 管路参数优化参考",
        "- 关联对象：C011 管路总成",
        "- 关联阶段：设计阶段",
        "- 建议反馈子系统：设计子系统",
        "- 说明：系统识别该质量反馈与管路总成设计参数相关，输出设计阶段参数复核建议，供下游管路参数优化模块读取。",
        f"- 下游读取变量：{'、'.join(str(item) for item in variables)}",
        "",
        "参数复核建议表：",
        md_table(
            ["参数名称", "参数类别", "优先级", "复核类型", "复核原因", "优化提示"],
            pipe_reference_rows(reference),
        ),
        "",
    ]


def render_markdown(report: dict[str, Any]) -> str:
    meta = report["report_meta"]
    fb = report["feedback_summary"]
    target = report["target_object"]
    rca = report["rca_summary"]
    graph = report["graph_comparison"]
    stage = report["stage_attribution"]
    subsystem = report["subsystem_feedback"]
    lines = [
        "# 全生命周期质量追溯分析报告",
        "",
        "## 1. 报告基本信息",
        f"- domain：{meta['domain']}",
        f"- case_id：{meta['case_id']}",
        f"- feedback_id：{meta['feedback_id']}",
        f"- report_time：{meta['report_time']}",
        f"- model_name：{meta['model_name']}",
        f"- top_k：{meta['top_k']}",
        f"- reasoning_json：{meta['source_files']['reasoning_json']}",
        f"- original_graph_json：{meta['source_files']['original_graph_json']}",
        f"- fault_enhanced_graph_json：{meta['source_files']['fault_enhanced_graph_json']}",
        "",
        "## 2. 质量反馈概述",
        f"- 故障类型：{fb['fault_type']}",
        f"- 故障位置：{fb['fault_position']}",
        f"- 严重程度：{fb['fault_severity']}",
        f"- 诊断置信度：{fb['diagnosis_confidence']}",
        f"- RCA置信度：{fb['rca_confidence']}",
        f"- 目标对象：{target['object_id']}（{target['object_type']}，{target['object_name']}）",
        "",
        "## 3. 上游诊断/RCA结果",
    ]
    if meta["domain"] == "hydraulic":
        lines.extend(
            [
                f"- 根因传感器：{rca.get('root_sensor', '')}（{rca.get('root_sensor_confidence', '')}）",
                f"- 次因传感器：{rca.get('secondary_sensor', '')}（{rca.get('secondary_sensor_confidence', '')}）",
                f"- 根因部件：{rca.get('root_component_code', '')} {rca.get('root_component_name', '')}",
                f"- 部件置信度：{rca.get('component_confidence', '')}",
                f"- 传播说明：{rca.get('propagation_note', '未发现明确跨部件传播描述。')}",
                "",
                "### 12部件诊断排序 Top5",
                md_table(["rank", "component_code", "component_name", "fault_type", "severity", "judgement", "fault_level"], [[r.get("rank_no"), r.get("component_code"), r.get("component_name"), r.get("fault_type"), r.get("severity"), r.get("judgement"), r.get("fault_level")] for r in rca.get("component_diagnosis_top5", [])]),
                "",
                "### 子类型概率 Top5",
                md_table(["component_code", "subtype_id", "subtype_name", "fault_part", "probability"], [[r.get("component_code"), r.get("subtype_id"), r.get("subtype_name"), r.get("fault_part"), r.get("probability")] for r in rca.get("subtype_probability_top5", [])]),
                "",
                "### 传感器异常能量 Top5",
                md_table(["sensor", "energy"], [[k, v] for k, v in rca.get("sensor_energy_top5", {}).items()]),
            ]
        )
    else:
        lines.extend([f"- RCA文本：{rca.get('rca_text', '')}", f"- 故障演化文本：{rca.get('fault_evolution_text', '')}"])
    lines.extend(
        [
            "",
            "## 4. 原始生命周期知识图谱概述",
            md_table(
                ["graph", "node_count", "edge_count", "包含生命周期阶段"],
                [["original_graph", graph["original_graph"]["node_count"], graph["original_graph"]["edge_count"], "、".join(graph["original_graph"]["stages"])]],
            ),
            "",
            "目标对象生命周期节点摘要：",
            md_table(["relation", "node_id", "node_type", "stage", "node_name"], [[r["relation_name"], r["node_id"], r["node_type"], r["stage_name"], r["node_name"]] for r in graph["original_graph"]["target_lifecycle_nodes"][:12]]),
            "",
            "## 5. 故障增强知识图谱概述",
            md_table(
                ["graph", "node_count", "edge_count", "added_node_count", "added_edge_count"],
                [["fault_enhanced_graph", graph["fault_enhanced_graph"]["node_count"], graph["fault_enhanced_graph"]["edge_count"], graph["fault_enhanced_graph"]["added_node_count"], graph["fault_enhanced_graph"]["added_edge_count"]]],
            ),
            f"\n{graph['difference_description']}",
            "",
            *reason_ranking_markdown_lines(report),
            "",
            "## 7. 生命周期阶段归因",
            f"- 首要阶段：{stage['primary_stage_name']}",
            f"- 反馈目标子系统：{stage['feedback_target_subsystem']}",
            f"- 说明：{stage['primary_stage_explanation']}",
            "",
            md_table(["stage", "stage_name", "score"], [[r["stage"], r["stage_name"], r["score"]] for r in stage["stage_rank"]]),
            *(
                [
                    "",
                    "### 重点关注阶段 Top3",
                    md_table(
                        ["rank", "stage", "stage_name", "score", "feedback_subsystem"],
                        [[r["rank"], r["stage"], r["stage_name"], r["score"], r["feedback_subsystem"]] for r in stage.get("focus_stage_top3", [])],
                    ),
                ]
                if stage.get("focus_stage_top3")
                else []
            ),
            "",
            "## 8. 关键证据链",
        ]
    )
    for chain in report["evidence_chains"]:
        lines.append(f"- {' -> '.join(chain['chain'])}：{chain['description']}")
    lines.extend(
        [
            "",
            "## 9. 子系统反馈建议",
            f"- 目标子系统：{subsystem['target_subsystem']}",
            f"- 首要反馈子系统：{subsystem.get('primary_feedback_subsystem', subsystem['target_subsystem'])}",
            *(
                [f"- 重点反馈子系统 Top3：{'、'.join(subsystem.get('feedback_subsystem_top3', []))}"]
                if subsystem.get("feedback_subsystem_top3")
                else []
            ),
            f"- 建议：{subsystem['suggestion']}",
            "",
            "## 10. 追溯结论",
            report["final_conclusion"],
            "",
        ]
    )
    if subsystem.get("review_subsystem_candidates"):
        lines.extend(
            [
                "",
                "### 生命周期阶段复核子系统参考",
                md_table(
                    ["stage", "stage_name", "score", "feedback_subsystem"],
                    [[r["stage"], r["stage_name"], r["score"], r["feedback_subsystem"]] for r in subsystem.get("review_subsystem_candidates", [])],
                ),
                "",
            ]
        )
    lines.extend(pipe_reference_markdown_lines(report))
    return sanitize_text("\n".join(lines))


def add_pipe_reference_docx_section(doc: Any, report: dict[str, Any]) -> None:
    if not pipe_design_reference_enabled(report):
        return
    reference = report["pipe_design_reference"]
    variables = reference.get("optimization_variables", [])
    if not isinstance(variables, list):
        variables = []

    doc.add_heading("10. 管路参数优化参考", level=1)
    doc.add_paragraph("关联对象：C011 管路总成")
    doc.add_paragraph("关联阶段：设计阶段")
    doc.add_paragraph("建议反馈子系统：设计子系统")
    doc.add_paragraph("说明：系统识别该质量反馈与管路总成设计参数相关，输出设计阶段参数复核建议，供下游管路参数优化模块读取。")
    doc.add_paragraph(f"下游读取变量：{'、'.join(str(item) for item in variables)}")

    table = doc.add_table(rows=1, cols=6)
    for i, heading in enumerate(["参数名称", "参数类别", "优先级", "复核类型", "复核原因", "优化提示"]):
        table.rows[0].cells[i].text = heading
    for row in pipe_reference_rows(reference):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = sanitize_text(value)


def add_reason_ranking_docx_section(doc: Any, report: dict[str, Any]) -> None:
    if pipe_design_reference_enabled(report) and isinstance(report.get("reasonList"), list):
        doc.add_heading("6. 最终溯源原因表", level=1)
        doc.add_paragraph("综合置信度综合考虑模型推理结果、图谱路径证据、反馈对象与生命周期阶段之间的关联关系以及管路设计参数证据。")
        table = doc.add_table(rows=1, cols=7)
        for i, heading in enumerate(["排名", "原因阶段/类型", "原因名称", "关联对象", "置信度", "关键证据", "建议措施"]):
            table.rows[0].cells[i].text = heading
        for row in reason_list_rows(report):
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = sanitize_text(value)
        return

    doc.add_heading("5. TransH Top6疑似原因分析", level=1)
    table = doc.add_table(rows=1, cols=7)
    for i, h in enumerate(["rank", "candidate_id", "type", "stage", "final", "transh", "scope"]):
        table.rows[0].cells[i].text = h
    for c in report["top6_reason_analysis"]:
        cells = table.add_row().cells
        for i, value in enumerate([c["rank"], c["candidate_id"], c["candidate_type"], c["candidate_stage_name"], c["final_score"], c["transh_score"], c["candidate_scope"]]):
            cells[i].text = sanitize_text(value)


def write_docx(report: dict[str, Any], path: Path) -> bool:
    try:
        from docx import Document
    except Exception as exc:
        print(f"python-docx 不可用，跳过 docx 输出：{exc}")
        return False
    doc = Document()
    doc.add_heading("全生命周期质量追溯分析报告", level=0)
    meta = report["report_meta"]
    fb = report["feedback_summary"]
    target = report["target_object"]
    doc.add_heading("1. 报告基本信息", level=1)
    doc.add_paragraph(f"domain：{meta['domain']}；case_id：{meta['case_id']}；feedback_id：{meta['feedback_id']}；report_time：{meta['report_time']}；model_name：TransH；top_k：6。")
    doc.add_heading("2. 质量反馈概述", level=1)
    doc.add_paragraph(f"故障类型：{fb['fault_type']}；故障位置：{fb['fault_position']}；严重程度：{fb['fault_severity']}；目标对象：{target['object_id']}（{target['object_name']}）。")
    doc.add_heading("3. 上游诊断/RCA结果", level=1)
    doc.add_paragraph(render_rca_paragraph(report))
    doc.add_heading("4. 图谱概述", level=1)
    graph = report["graph_comparison"]
    table = doc.add_table(rows=1, cols=5)
    for i, h in enumerate(["graph", "nodes", "edges", "added_nodes", "added_edges"]):
        table.rows[0].cells[i].text = h
    for row in [["original_graph", graph["original_graph"]["node_count"], graph["original_graph"]["edge_count"], "-", "-"], ["fault_enhanced_graph", graph["fault_enhanced_graph"]["node_count"], graph["fault_enhanced_graph"]["edge_count"], graph["fault_enhanced_graph"]["added_node_count"], graph["fault_enhanced_graph"]["added_edge_count"]]]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = sanitize_text(value)
    add_reason_ranking_docx_section(doc, report)
    doc.add_heading("6. 生命周期阶段归因", level=1)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["stage", "stage_name", "score"]):
        table.rows[0].cells[i].text = h
    for r in report["stage_attribution"]["stage_rank"]:
        cells = table.add_row().cells
        for i, value in enumerate([r["stage"], r["stage_name"], r["score"]]):
            cells[i].text = sanitize_text(value)
    doc.add_paragraph(report["stage_attribution"]["primary_stage_explanation"])
    focus_stage_top3 = report["stage_attribution"].get("focus_stage_top3", [])
    if focus_stage_top3:
        doc.add_heading("重点关注阶段 Top3", level=2)
        table = doc.add_table(rows=1, cols=5)
        for i, h in enumerate(["rank", "stage", "stage_name", "score", "feedback_subsystem"]):
            table.rows[0].cells[i].text = h
        for r in focus_stage_top3:
            cells = table.add_row().cells
            for i, value in enumerate([r["rank"], r["stage"], r["stage_name"], r["score"], r["feedback_subsystem"]]):
                cells[i].text = sanitize_text(value)
    doc.add_heading("7. 关键证据链", level=1)
    for chain in report["evidence_chains"]:
        doc.add_paragraph(f"{' -> '.join(chain['chain'])}：{chain['description']}", style=None)
    doc.add_heading("8. 子系统反馈建议", level=1)
    subsystem = report["subsystem_feedback"]
    doc.add_paragraph(f"首要反馈子系统：{subsystem.get('primary_feedback_subsystem', subsystem.get('target_subsystem', ''))}")
    subsystem_top3 = subsystem.get("feedback_subsystem_top3", [])
    if subsystem_top3:
        doc.add_paragraph(f"重点反馈子系统 Top3：{'、'.join(subsystem_top3)}")
    doc.add_paragraph(report["subsystem_feedback"]["suggestion"])
    doc.add_heading("9. 追溯结论", level=1)
    doc.add_paragraph(report["final_conclusion"])
    add_pipe_reference_docx_section(doc, report)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return True


def render_rca_paragraph(report: dict[str, Any]) -> str:
    if report["report_meta"]["domain"] == "hydraulic":
        r = report["rca_summary"]
        return sanitize_text(f"根因传感器：{r.get('root_sensor', '')}；次因传感器：{r.get('secondary_sensor', '')}；根因部件：{r.get('root_component_code', '')} {r.get('root_component_name', '')}。{r.get('propagation_note', '')}")
    r = report["rca_summary"]
    return sanitize_text(f"RCA文本：{r.get('rca_text', '')} 故障演化文本：{r.get('fault_evolution_text', '')}")


def build_report_for_case(
    domain: str,
    case_id_arg: str,
    reasoning: dict[str, Any] | None = None,
    save_files: bool = True,
    generate_docx: bool = True,
) -> dict[str, Any]:
    report_dir = ROOT / "outputs" / domain / "reports"
    if reasoning is not None and save_files:
        safe_case_id = str(reasoning.get("case_id", normalize_case_id_alias(case_id_arg))).replace("/", "_").replace("\\", "_")
        report_dir.mkdir(parents=True, exist_ok=True)
        reasoning_path = report_dir / f"reasoning_{safe_case_id}.json"
        reasoning_path.write_text(json.dumps(reasoning, ensure_ascii=False, indent=2), encoding="utf-8")

    paths = resolve_paths(domain, case_id_arg)
    report = build_report(domain, paths)
    case_id = report["report_meta"]["case_id"]
    json_path = report_dir / f"trace_report_{case_id}.json"
    md_path = report_dir / f"trace_report_{case_id}.md"
    docx_path = report_dir / f"trace_report_{case_id}.docx"
    markdown = render_markdown(report)
    docx_ok = False
    if save_files:
        report_dir.mkdir(parents=True, exist_ok=True)
        json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        md_path.write_text(markdown, encoding="utf-8")
        if generate_docx:
            docx_ok = write_docx(report, docx_path)
    return {
        "domain": domain,
        "case_id": case_id,
        "report_json": report,
        "report_markdown": markdown,
        "artifacts": {
            "trace_report_json": str(json_path),
            "trace_report_md": str(md_path),
            "trace_report_docx": str(docx_path),
            "report_json": str(json_path),
            "report_md": str(md_path),
            "report_docx": str(docx_path),
        },
        "docx_generated": docx_ok,
    }


def main() -> None:
    args = parse_args()
    result = build_report_for_case(args.domain, args.case_id, save_files=True, generate_docx=not args.no_docx)
    json_path = Path(result["artifacts"]["trace_report_json"])
    md_path = Path(result["artifacts"]["trace_report_md"])
    docx_path = Path(result["artifacts"]["trace_report_docx"])
    print(f"JSON报告路径: {json_path}")
    print(f"Markdown报告路径: {md_path}")
    if result["docx_generated"]:
        print(f"Word报告路径: {docx_path}")


    if args.print_json:
        print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
